"""Merge parsed RCAs, auto scores and judgment reviews into the master workbook.

Output: out/RCA Master Database.xlsx with sheets Scorecard, Issue Summary, RCA Log, Actions, Recurrence.
CI managers can override the suggested grade on the Scorecard; overrides are read back from the existing
workbook (and data/overrides.json) on every rebuild so they are never lost.
"""
import collections
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from paths import HOME as ROOT, MASTER  # noqa: E402
JUDGE_FIELDS = ["problem_specific", "problem_not_cause", "where_else", "refocus", "standards_gaps_actioned",
                "maintenance_present", "gemba_evidence", "gemba_confirmed", "rc_why", "rc_controllable",
                "rc_chain", "lt_action_targets_rc", "actions_map", "lt_durable"]
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
BAND_FILL = {"Solid": "E3F1E6", "Coach": "FFF4D6", "Redo": "FBE3E1"}
OUT = MASTER
OVERRIDES = ROOT / "data" / "overrides.json"
OVERRIDE_COLS = ["CI grade override", "Override reason", "Overridden by"]

# Issue types, in the order CI managers should care about them. Each matches flag text from
# score_auto.py (mechanical) or the judgment review's hard_flags / recurrence verdicts.
ISSUE_TYPES = [
    ("Repeat failure", lambda f: f.startswith("Recurrence of")),
    ("Root cause is a symptom", lambda f: "describes what failed" in f),
    ("No verified root cause", lambda f: f.startswith("No verified root cause") and "describes what failed" not in f),
    ("No long-term fix for standards gap", lambda f: f.startswith("No long-term action against")),
    ("One-person brainstorm", lambda f: f.startswith("One-person brainstorm")),
    ("Gemba skipped", lambda f: f.startswith("Gemba section left blank")),
    ("No actions", lambda f: f.startswith("No actions listed")),
    ("Action missing owner or due date", lambda f: "missing an owner or due date" in f),
    ("Closed without completion date", lambda f: "resolved with no completion date" in f),
]


def issue_types(flags):
    return [name for name, match in ISSUE_TYPES if any(match(f) for f in flags)]


def load_overrides():
    """Overrides typed into the existing workbook win over the saved copy."""
    saved = json.loads(OVERRIDES.read_text(encoding="utf-8")) if OVERRIDES.exists() else {}
    if OUT.exists():
        ws = load_workbook(OUT)["Scorecard"]
        hdr = [c.value for c in ws[1]]
        if "ID" in hdr and OVERRIDE_COLS[0] in hdr:
            ci = {h: hdr.index(h) for h in ["ID"] + OVERRIDE_COLS}
            for row in ws.iter_rows(min_row=2, values_only=True):
                rid = row[ci["ID"]]
                vals = {h: (row[ci[h]] or "") for h in OVERRIDE_COLS}
                if rid and any(vals.values()):
                    saved[rid] = vals
                elif rid in saved and not any(vals.values()):
                    del saved[rid]  # cleared in the spreadsheet
    OVERRIDES.write_text(encoding="utf-8", data=json.dumps(saved, indent=1))
    return saved


def band(score, flags):
    if score is None:
        return ""
    if flags or score < 70:
        return "Redo"
    return "Solid" if score >= 85 else "Coach"


def load():
    rcas = json.loads((ROOT / "data" / "rcas.json").read_text(encoding="utf-8"))
    auto = json.loads((ROOT / "data" / "auto_scores.json").read_text(encoding="utf-8"))
    judged = {}
    for f in sorted((ROOT / "data").glob("judge_output_*.json")):
        for j in json.loads(f.read_text(encoding="utf-8")):
            judged[j["id"]] = j
    return rcas, auto, judged


def dedupe_key(r):
    text = re.sub(r"[^a-z0-9]", "", (r["initial_problem"] + r["w2h"].get("what", "")).lower())
    return (r["event_date"], text) if r["event_date"] and text else None


def merge(rcas, auto, judged):
    # The same RCA is sometimes pasted into two line workbooks; keep the first, note the other.
    seen = {}
    # Main line workbooks win over stray copies (codes like "ML3-BeltAdj")
    for r in sorted(rcas, key=lambda r: "-" in r["code"]):
        k = dedupe_key(r)
        if k in seen:
            r["duplicate_of"] = seen[k]["id"]
            seen[k].setdefault("also_filed_in", []).append(r["line"])
        elif k:
            seen[k] = r
    rows = []
    for r in rcas:
        if r.get("duplicate_of"):
            continue
        a, j = auto[r["id"]], judged.get(r["id"])
        flags = list(a["flags"])
        score = None
        if j:
            score = a["auto_points"] + sum(int(j.get(k, 0)) for k in JUDGE_FIELDS)
            flags += j.get("hard_flags", [])
            for rc in j.get("recurrence", []):
                if rc["verdict"] == "same_failure":
                    flags.append(f"Recurrence of {rc['prior_id'].split('::', 1)[1]} — prior fix did not hold.")
        notes = [f"Also filed in {', '.join(r['also_filed_in'])} workbook."] if r.get("also_filed_in") else []
        rows.append({"r": r, "a": a, "j": j, "score": score, "flags": flags + notes, "band": band(score, flags),
                     "issues": issue_types(flags)})
    return rows


def write_sheet(ws, headers, data, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in data:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    rcas, auto, judged = load()
    rows = merge(rcas, auto, judged)
    rows.sort(key=lambda x: x["r"]["event_date"] or "", reverse=True)
    wb = Workbook()

    overrides = load_overrides()
    sc = wb.active
    sc.title = "Scorecard"
    reviewed = [x for x in rows if x["score"] is not None]
    issue_names = [n for n, _ in ISSUE_TYPES]
    headers = (["ID", "Event date", "Plant", "Line", "RCA", "Author", "Score", "Suggested grade"] + OVERRIDE_COLS
               + ["Final grade", "Issue types"] + issue_names + ["Coaching", "Summary", "Flag detail"])
    data = []
    for n, x in enumerate(reviewed, start=2):
        ov = overrides.get(x["r"]["id"], {})
        data.append([x["r"]["id"], x["r"]["event_date"], x["r"]["plant"], x["r"]["line"], x["r"]["sheet"],
                     x["r"].get("author", ""), x["score"], x["band"]] + [ov.get(h, "") for h in OVERRIDE_COLS]
                    + [f'=IF(I{n}<>"",I{n},H{n})', ", ".join(x["issues"])]
                    + ["X" if t in x["issues"] else "" for t in issue_names]
                    + ["\n".join(x["j"].get("coaching", [])), x["j"].get("summary", ""), "\n".join(x["flags"])])
    write_sheet(sc, headers, data, [22, 11, 6, 16, 26, 16, 7, 10, 11, 28, 14, 10, 34] + [9] * len(issue_names) + [60, 50, 45])
    last = len(data) + 1
    dv = DataValidation(type="list", formula1='"Solid,Coach,Redo"', allow_blank=True,
                        promptTitle="CI grade override", prompt="Pick a grade to override the suggested one; leave blank to accept it.")
    sc.add_data_validation(dv)
    dv.add(f"I2:I{max(last, 2)}")
    for col in ("H", "I", "L"):
        for g, color in BAND_FILL.items():
            sc.conditional_formatting.add(f"{col}2:{col}{max(last, 2)}",
                                          CellIsRule(operator="equal", formula=[f'"{g}"'], fill=PatternFill("solid", fgColor=color)))
    for c in sc[1][8:11]:
        c.fill = PatternFill("solid", fgColor="5A4A00")
    for row in sc.iter_rows(min_row=2, min_col=14, max_col=13 + len(issue_names)):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="top")

    summ = wb.create_sheet("Issue Summary")
    lines = sorted({x["r"]["line"] for x in reviewed})
    summ_rows = []
    for t in issue_names:
        hit = [x for x in reviewed if t in x["issues"]]
        by_line = collections.Counter(x["r"]["line"] for x in hit)
        summ_rows.append([t, len(hit), f"{len(hit) / max(len(reviewed), 1):.0%}"] + [by_line.get(l, 0) or "" for l in lines])
    summ_rows.append(["RCAs reviewed", len(reviewed), ""] + [sum(1 for x in reviewed if x["r"]["line"] == l) for l in lines])
    write_sheet(summ, ["Issue type", "RCAs flagged", "Share of reviewed"] + lines, summ_rows, [34, 10, 10] + [11] * len(lines))
    summ.cell(row=len(summ_rows) + 1, column=1).font = Font(bold=True)

    log = wb.create_sheet("RCA Log")
    write_sheet(log, ["ID", "Event date", "Plant", "Line", "RCA", "Author", "Problem statement", "Refocused statement",
                      "Gemba root causes", "Brainstorm attendees", "Auto points (of 41)", "Rigor score", "Grade",
                      "Open actions", "Flags", "Source workbook"],
                [[x["r"]["id"], x["r"]["event_date"], x["r"]["plant"], x["r"]["line"], x["r"]["sheet"],
                  x["r"].get("author", ""), x["r"]["initial_problem"], x["r"]["refocus"], "\n".join(x["r"]["gemba"]),
                  ", ".join(x["r"]["brainstorm_attendees"]), x["a"]["auto_points"], x["score"], x["band"],
                  x["a"]["open_actions"], "\n".join(x["flags"]), x["r"]["source_file"].rsplit("_p", 1)[0]]
                 for x in rows],
                [30, 11, 6, 18, 28, 18, 45, 50, 40, 25, 9, 8, 8, 8, 40, 10])

    acts = wb.create_sheet("Actions")
    act_rows = []
    today = date.today().isoformat()
    for x in rows:
        for kind in ("short_term", "long_term"):
            for a in x["r"][kind]:
                if not a["action"] or a["action"].upper() == "N/A":
                    continue
                act_rows.append([x["r"]["id"], x["r"]["event_date"], x["r"]["line"], "Short term" if kind == "short_term" else "Long term",
                                 a["action"], a["owner"], a["due"], a["completed"], a["resolved"]])
    write_sheet(acts, ["RCA ID", "Event date", "Line", "Type", "Action", "Owner", "Due", "Completed", "Resolved"],
                act_rows, [30, 11, 18, 10, 60, 18, 11, 11, 9])

    rec = wb.create_sheet("Recurrence")
    rec_rows = []
    for x in rows:
        if x["j"]:
            for rc in x["j"].get("recurrence", []):
                if rc["verdict"] in ("same_failure", "related"):
                    rec_rows.append([x["r"]["event_date"], x["r"]["line"], x["r"]["sheet"], rc["prior_id"].split("::", 1)[1],
                                     "Same failure" if rc["verdict"] == "same_failure" else "Related", rc["note"]])
    write_sheet(rec, ["Event date", "Line", "New RCA", "Earlier RCA", "Verdict", "Note"], rec_rows, [11, 18, 30, 30, 12, 70])

    out = OUT
    out.parent.mkdir(exist_ok=True)
    try:
        wb.save(out)
    except PermissionError:
        # Windows locks a workbook that is open in desktop Excel
        sys.exit(f"Could not save {out}: it is open in Excel on this computer. Close it and rerun build_db.py.")
    print(f"{len(rows)} RCAs, {len(reviewed)} reviewed, {len(act_rows)} actions, {len(rec_rows)} recurrence links -> {out}")


if __name__ == "__main__":
    main()
